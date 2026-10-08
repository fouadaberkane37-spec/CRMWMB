import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

STATUSES = ["Booked", "Pending", "Cancelled"]

NAVY   = "1E293B"
INDIGO = "4F46E5"
STRIPE = "F8FAFC"
INPUT  = "FFFBEB"
GRID   = "CBD5E1"

COLUMNS = [
    # (header, width, is_caller_input)
    ("#",               5,  False),
    ("Client",          26, False),
    ("Phone",           16, False),
    ("Address",         36, False),
    ("Last Service",    14, False),
    ("Services Done",   30, False),
    ("Last Paid ($)",   13, False),
    ("Status",          14, True),
    ("New Date",        14, True),
    ("Called On",       14, True),
    ("Notes",           40, True),
]

TITLE_ROW   = 1
SUBTITLE_ROW = 2
SUMMARY_LABEL_ROW = 4
SUMMARY_VALUE_ROW = 5
HEADER_ROW  = 7
FIRST_DATA  = HEADER_ROW + 1


def _fmt_phone(raw: str) -> str:
    digits = "".join(ch for ch in (raw or "") if ch.isdigit())
    if len(digits) == 11 and digits.startswith("1"):
        digits = digits[1:]
    if len(digits) == 10:
        return f"({digits[:3]}) {digits[3:6]}-{digits[6:]}"
    return raw or ""


def build_reactivation_workbook(rows: list[dict]) -> bytes:
    """rows: dicts with keys name, phone, address, last_service (datetime|None), services, paid."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Call Sheet"
    ws.sheet_view.showGridLines = False

    last_col = len(COLUMNS)
    last_letter = get_column_letter(last_col)
    thin = Side(style="thin", color=GRID)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # Title
    ws.merge_cells(start_row=TITLE_ROW, start_column=1, end_row=TITLE_ROW, end_column=last_col)
    t = ws.cell(row=TITLE_ROW, column=1, value="WMB — Reactivation Call Sheet")
    t.font = Font(size=18, bold=True, color="FFFFFF")
    t.fill = PatternFill("solid", fgColor=NAVY)
    t.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[TITLE_ROW].height = 34

    ws.merge_cells(start_row=SUBTITLE_ROW, start_column=1, end_row=SUBTITLE_ROW, end_column=last_col)
    s = ws.cell(
        row=SUBTITLE_ROW, column=1,
        value=f"Past clients — job completed & paid  •  {len(rows)} to call  •  generated {datetime.now().strftime('%Y-%m-%d')}",
    )
    s.font = Font(size=10, italic=True, color="64748B")
    s.alignment = Alignment(horizontal="left", indent=1)

    # Live summary (COUNTIF over the Status column)
    status_col = next(i for i, c in enumerate(COLUMNS, 1) if c[0] == "Status")
    status_letter = get_column_letter(status_col)
    last_data = FIRST_DATA + max(len(rows), 1) - 1
    status_range = f"${status_letter}${FIRST_DATA}:${status_letter}${last_data}"

    summary = [
        ("Booked",    f'=COUNTIF({status_range},"Booked")',    "059669"),
        ("Pending",   f'=COUNTIF({status_range},"Pending")',   "D97706"),
        ("Cancelled", f'=COUNTIF({status_range},"Cancelled")', "DC2626"),
        ("Not called yet", f'=ROWS({status_range})-COUNTA({status_range})', "475569"),
    ]
    col = 2
    for label, formula, color in summary:
        lc = ws.cell(row=SUMMARY_LABEL_ROW, column=col, value=label)
        lc.font = Font(size=9, bold=True, color="FFFFFF")
        lc.fill = PatternFill("solid", fgColor=color)
        lc.alignment = Alignment(horizontal="center")
        vc = ws.cell(row=SUMMARY_VALUE_ROW, column=col, value=formula)
        vc.font = Font(size=16, bold=True, color=color)
        vc.alignment = Alignment(horizontal="center")
        vc.border = border
        col += 1
    ws.row_dimensions[SUMMARY_VALUE_ROW].height = 26

    # Header
    for i, (header, width, is_input) in enumerate(COLUMNS, 1):
        c = ws.cell(row=HEADER_ROW, column=i, value=header)
        c.font = Font(bold=True, color="FFFFFF", size=11)
        c.fill = PatternFill("solid", fgColor=INDIGO if is_input else NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[HEADER_ROW].height = 24

    # Data
    for idx, r in enumerate(rows):
        row = FIRST_DATA + idx
        values = [
            idx + 1,
            r.get("name") or "",
            _fmt_phone(r.get("phone") or ""),
            r.get("address") or "",
            r.get("last_service").date() if r.get("last_service") else None,
            r.get("services") or "",
            float(r.get("paid") or 0),
            None, None, None, None,
        ]
        stripe = PatternFill("solid", fgColor=STRIPE) if idx % 2 else None
        for ci, (val, (_, _, is_input)) in enumerate(zip(values, COLUMNS), 1):
            c = ws.cell(row=row, column=ci, value=val)
            c.border = border
            c.alignment = Alignment(vertical="center", wrap_text=ci in (4, 6, 11))
            if is_input:
                c.fill = PatternFill("solid", fgColor=INPUT)
            elif stripe:
                c.fill = stripe
        ws.cell(row=row, column=1).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=row, column=2).font = Font(bold=True)
        ws.cell(row=row, column=5).number_format = "yyyy-mm-dd"
        ws.cell(row=row, column=7).number_format = '"$"#,##0.00'
        ws.cell(row=row, column=9).number_format = "yyyy-mm-dd"
        ws.cell(row=row, column=10).number_format = "yyyy-mm-dd"
        ws.cell(row=row, column=status_col).alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[row].height = 30

    # Status dropdown + colour coding
    dv = DataValidation(type="list", formula1=f'"{",".join(STATUSES)}"', allow_blank=True)
    dv.error = "Pick Booked, Pending or Cancelled"
    dv.errorTitle = "Invalid status"
    dv.prompt = "Booked / Pending / Cancelled"
    dv.promptTitle = "Call outcome"
    ws.add_data_validation(dv)
    dv.add(f"{status_letter}{FIRST_DATA}:{status_letter}{last_data}")

    date_dv = DataValidation(type="date", operator="greaterThan", formula1="DATE(2020,1,1)", allow_blank=True)
    date_dv.error = "Enter a date (YYYY-MM-DD)"
    ws.add_data_validation(date_dv)
    for letter in ("I", "J"):
        date_dv.add(f"{letter}{FIRST_DATA}:{letter}{last_data}")

    data_range = f"A{FIRST_DATA}:{last_letter}{last_data}"
    first_status = f"${status_letter}{FIRST_DATA}"
    for value, bg, fg in (
        ("Booked",    "D1FAE5", "065F46"),
        ("Pending",   "FEF3C7", "92400E"),
        ("Cancelled", "FEE2E2", "991B1B"),
    ):
        ws.conditional_formatting.add(
            data_range,
            FormulaRule(
                formula=[f'{first_status}="{value}"'],
                fill=PatternFill("solid", fgColor=bg, bgColor=bg),
                font=Font(color=fg),
            ),
        )

    ws.freeze_panes = ws.cell(row=FIRST_DATA, column=3)
    ws.auto_filter.ref = f"A{HEADER_ROW}:{last_letter}{last_data}"

    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{HEADER_ROW}:{HEADER_ROW}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
